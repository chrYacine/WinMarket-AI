"""QA campaign 2026-09-13 — closing three specific gaps identified while
mapping the existing B02/B03 suite against QA01-QA30:

QA12: extraction provenance — a DOCX table and a multi-page PDF must
produce chunks whose page_number/section are actually usable, not just
accepted as optional columns.
QA26: anonymous (no session at all) requests to the new knowledge routes
are refused, not just role-insufficient ones.
QA28: changing an account's documentation/capacity after an analysis was
already run must never rewrite that analysis's already-persisted result.
"""
from __future__ import annotations

import io
import re

import uuid as uuid_module

from docx import Document as DocxDocument

from tests.conftest import default_org_id, make_active_starter_user

# B04-T0: filesystem isolation and the real-network block are now autouse
# at tests/conftest.py level for the whole suite — see
# _b04_isolated_filesystem_roots and _b04_block_real_network there.


def _csrf_from(html: str) -> str:
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert match
    return match.group(1)


def _login(client, email, password="Sup3rSecret!"):
    r = client.get("/login")
    csrf = _csrf_from(r.text)
    client.post("/login", data={"email": email, "password": password, "next": "/app", "csrf_token": csrf})
    return csrf


def _upload(client, filename: str, content: bytes, csrf: str | None = None, content_type: str = "text/plain"):
    return client.post(
        "/api/knowledge/documents",
        files={"file": (filename, io.BytesIO(content), content_type)},
        headers={"X-CSRF-Token": csrf} if csrf is not None else {},
    )


# ---------------------------------------------------------------------------
# QA12 — extraction preserves usable page/table provenance
# ---------------------------------------------------------------------------

def test_docx_table_is_extracted_as_a_located_chunk(client, db):
    from src.web.database.repositories import knowledge as knowledge_repo

    user = make_active_starter_user(db, "qa12docx@example.com")
    org_id = default_org_id(db, user)
    csrf = _login(client, "qa12docx@example.com")

    doc = DocxDocument()
    doc.add_paragraph("Texte d'introduction PREUVE_QA12_PARAGRAPH.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Colonne A"
    table.rows[0].cells[1].text = "Colonne B"
    table.rows[1].cells[0].text = "PREUVE_QA12_TABLE"
    table.rows[1].cells[1].text = "valeur"
    buf = io.BytesIO()
    doc.save(buf)

    r = client.post("/api/knowledge/documents", files={"file": ("with_table.docx", buf.getvalue(),
                     "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                     headers={"X-CSRF-Token": csrf})
    assert r.status_code == 201
    doc_id = r.json()["document"]["id"]

    document = knowledge_repo.get_document_for_owner(db, document_id=uuid_module.UUID(doc_id), organization_id=org_id, owner_user_id=user.id)
    version = document.active_version
    assert version is not None and version.extraction_status == "ready"
    assert any(c.section and "Tableau" in c.section and "PREUVE_QA12_TABLE" in c.content for c in version.chunks)
    assert any("PREUVE_QA12_PARAGRAPH" in c.content and c.section is None for c in version.chunks)

    search = client.get("/api/knowledge/search?q=PREUVE_QA12_TABLE").json()
    assert any("PREUVE_QA12_TABLE" in r["excerpt"] for r in search["results"])


def test_pdf_chunk_carries_correct_page_number(client, db):
    from src.web.database.repositories import knowledge as knowledge_repo

    user = make_active_starter_user(db, "qa12pdf@example.com")
    org_id = default_org_id(db, user)
    csrf = _login(client, "qa12pdf@example.com")

    import fitz
    pdf_doc = fitz.open()
    page1 = pdf_doc.new_page()
    page1.insert_text((72, 72), "PREUVE_QA12_PAGE_UN sur la premiere page.")
    page2 = pdf_doc.new_page()
    page2.insert_text((72, 72), "PREUVE_QA12_PAGE_DEUX sur la deuxieme page.")
    raw = pdf_doc.tobytes()
    pdf_doc.close()

    r = client.post("/api/knowledge/documents", files={"file": ("two_pages.pdf", io.BytesIO(raw), "application/pdf")},
                     headers={"X-CSRF-Token": csrf})
    assert r.status_code == 201
    doc_id = r.json()["document"]["id"]

    document = knowledge_repo.get_document_for_owner(db, document_id=uuid_module.UUID(doc_id), organization_id=org_id, owner_user_id=user.id)
    chunks = document.active_version.chunks
    page_for_marker_1 = next(c.page_number for c in chunks if "PREUVE_QA12_PAGE_UN" in c.content)
    page_for_marker_2 = next(c.page_number for c in chunks if "PREUVE_QA12_PAGE_DEUX" in c.content)
    assert page_for_marker_1 == 1
    assert page_for_marker_2 == 2


# ---------------------------------------------------------------------------
# QA26 — anonymous requests (no session at all) are refused
# ---------------------------------------------------------------------------

def test_anonymous_requests_to_knowledge_routes_are_refused(client, db):
    """No _login() call anywhere in this test — the TestClient carries no
    session cookie at all."""
    r_upload = _upload(client, "x.md", b"# X\n\ncontent")
    assert r_upload.status_code == 401
    assert b"content" not in r_upload.content

    r_list = client.get("/api/knowledge/documents")
    assert r_list.status_code == 401

    r_search = client.get("/api/knowledge/search?q=x")
    assert r_search.status_code == 401

    r_capacity = client.post("/api/capacity", json={"charge_globale_pct": 1, "nombre_projets_en_cours": 0, "projets_en_cours": [], "capacites_par_pole": {}})
    assert r_capacity.status_code == 401

    r_page = client.get("/app/base-connaissances", follow_redirects=False)
    assert r_page.status_code == 303  # HTML page: redirect to /login, never a bare 200
    assert r_page.headers["location"].startswith("/login")


# ---------------------------------------------------------------------------
# QA28 — an analysis's persisted result is never rewritten by a later
# documentation/capacity change
# ---------------------------------------------------------------------------

def test_analysis_result_is_immutable_after_later_documentation_and_capacity_changes(client, db, tmp_path, monkeypatch):
    from src.core import config
    from src.web import jobs as jobs_module

    output_dir = tmp_path / "outputs"
    output_dir.mkdir()
    monkeypatch.setattr(config, "OUTPUT_DIR", output_dir)
    monkeypatch.setattr(config, "LOCAL_STORAGE_PATH", tmp_path)
    monkeypatch.setattr(jobs_module, "ANALYSIS_FILES_DIR", output_dir)
    monkeypatch.setattr(config, "LLM_ENABLED", False)

    make_active_starter_user(db, "qa28@example.com")
    csrf = _login(client, "qa28@example.com")
    _upload(client, "before.md", b"# Before\n\nPREUVE_QA28_AVANT reference initiale.", csrf)

    valid_ao_text = (
        "Appel d'offres - Test provenance\n"
        "Acheteur : Collectivite Exemple\n"
        "Le prestataire realisera le projet et ses livrables.\n"
        "Budget : 200 000 euros. Date limite : 30/11/2026.\n"
        "Exigences : Python, FastAPI.\n"
    )
    r = client.post("/api/analyze", data={"mode": "paste", "text": valid_ao_text}, headers={"X-CSRF-Token": csrf})
    assert r.status_code == 200
    job_id = r.json()["job_id"]

    from src.web import jobs
    import time
    for _ in range(50):
        job = jobs.get_job(job_id)
        if job.status != "running":
            break
        time.sleep(0.1)
    assert job.status == "done", job.error

    result_before = job.result.model_dump()

    # Change documentation and capacity AFTER the analysis already ran.
    _upload(client, "after.md", b"# After\n\nPREUVE_QA28_APRES reference ajoutee apres coup.", csrf)
    client.post("/api/capacity", json={"charge_globale_pct": 99, "nombre_projets_en_cours": 5, "projets_en_cours": [], "capacites_par_pole": {}}, headers={"X-CSRF-Token": csrf})

    # The in-memory Job object (and its persisted JSON/DB row) must not
    # have been touched by either change — it is a snapshot, not a live view.
    job_after = jobs.get_job(job_id)
    assert job_after.result.model_dump() == result_before
    assert "PREUVE_QA28_APRES" not in str(job_after.result.model_dump())

    from src.web.database.repositories import analyses as analyses_repo
    from src.web.database.repositories import users as users_repo
    user = users_repo.get_by_email(db, "qa28@example.com")
    analysis = analyses_repo.get_by_job_id_for_user(db, job_id, user.id)
    assert analysis is not None
    assert "PREUVE_QA28_APRES" not in str(analysis.result_data)
