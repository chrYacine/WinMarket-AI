"""B19-T2 — document rendering decoupled from the analysis result.

Three confirmed defects:

1. A PDF/DOCX rendering failure ends the job at status="error"/
   document_generation_failed while the SQL analysis row and its snapshot
   are intact (B11-T1) — but the user had NO way to obtain the documents
   afterwards, and there was no per-document status at all ("PDF failed,
   DOCX fine" was indistinguishable from "both failed").
2. No regeneration action existed: once job.files was empty, the only
   recourse was re-running the whole analysis — re-paying for the LLM and,
   worse, risking a DIFFERENT result if the account's ScoringPolicy had
   changed since, exactly what B11-T1 exists to prevent.
3. Free-form AO/LLM text reached ReportLab's markup-parsing Paragraph
   unescaped, so an "&" or "<" in real business data could be swallowed or
   break the render.

Test groups, per the ticket:
A. Per-document status, reported separately; a physically deleted PDF
   reads "unavailable" while the DOCX stays "available".
B. The core scenario: rendering fails outright, then regeneration rebuilds
   the document from the snapshot alone — zero LLM calls (an adapter that
   raises on ANY access), every scoring field preserved field-by-field, a
   second call creating no duplicate row, and markup-shaped text rendered
   inert.
C. "INCOMPLET" survives regeneration verbatim.
D. Another authenticated user is refused, with the same shape as a
   genuinely missing analysis.
E. A snapshot missing "result", and one with a non-finite score, are both
   refused explicitly — never filled with placeholder content.
F. The escaping fix itself, plus the explicit confirmation that DOCX needs
   no equivalent.

Helpers are copied locally (this codebase's convention: test files never
import from one another). No real network, no real key, no commit.
"""
from __future__ import annotations

import re
import time

import pytest
from sqlalchemy import func, select

from src.web import document_rendering_service as drs
from src.web import jobs
from src.web.database.models import Analysis, AnalysisDocument
from tests.conftest import default_org_id, make_active_starter_user

VALID_WEIGHTS = {
    "Adequation expertise": 20, "References similaires": 15, "Disponibilite equipe": 10,
    "Rentabilite estimee": 10, "Faisabilite delai": 10, "Certifications requises": 10,
    "Complexite technique": 5, "Connaissance secteur": 5, "Potentiel commercial": 5,
    "Risque contractuel": 5, "Solidite client": 3, "Valeur strategique": 2,
}

# Fully configured rules → the engine reaches one of the three REAL
# verdicts. Leaving business_rules unset is the legal account state whose
# honest verdict is "INCOMPLET" (B06-T4) — group C needs exactly that.
COMPLETE_BUSINESS_RULES = {
    "budget_minimum_eur": 0, "max_charge_pct": 100,
    "max_unmastered_technologies": 999, "certification_penalty_score": 20,
}

VALID_AO_TEXT = (
    "Appel d'offres - Portail client\n"
    "Acheteur : Collectivite Exemple\n"
    "Budget : 250 000 euros. Date limite : 30/11/2026.\n"
    "Exigences : Python et Django.\n"
)


# ---------------------------------------------------------------------------
# Local helpers (copied, per this suite's no-cross-import convention).
# ---------------------------------------------------------------------------

@pytest.fixture(autouse=True)
def isolated_storage(tmp_path, monkeypatch):
    """config.OUTPUT_DIR, jobs.ANALYSIS_FILES_DIR and jobs.ANALYSES_DIR are
    all read once at import time from the REAL data directory, so conftest's
    autouse redirection does not reach the last two. Every one of them is
    pointed at a throwaway directory here — these tests delete generated
    files on purpose, which must never mean deleting anything under the
    project's own data/."""
    from src.core import config
    output_dir = tmp_path / "outputs"
    output_dir.mkdir()
    monkeypatch.setattr(config, "OUTPUT_DIR", output_dir)
    monkeypatch.setattr(config, "LOCAL_STORAGE_PATH", tmp_path)
    monkeypatch.setattr(jobs, "ANALYSIS_FILES_DIR", output_dir)
    monkeypatch.setattr(jobs, "ANALYSES_DIR", tmp_path / "historique" / "analyses")
    return output_dir


def _csrf_from(html: str) -> str:
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert match
    return match.group(1)


def _login(client, email, password="Sup3rSecret!"):
    r = client.get("/login")
    csrf = _csrf_from(r.text)
    client.post("/login", data={"email": email, "password": password, "next": "/app", "csrf_token": csrf})
    return csrf


def _configure_account(client, db, email, *, business_rules=None):
    """Returns (user, csrf) — csrf is the token to thread into any further
    mutating call in this session (B14-T1: X-CSRF-Token header)."""
    user = make_active_starter_user(db, email, scoring=False)
    csrf = _login(client, email)
    headers = {"X-CSRF-Token": csrf}
    client.post("/api/capacity", json={
        "charge_globale_pct": 40, "nombre_projets_en_cours": 1,
        "projets_en_cours": ["Projet test"], "capacites_par_pole": {"Software Engineering": 40},
    }, headers=headers)
    client.put("/api/scoring-config/profile", json={
        "raison_sociale": "ESN de test", "effectif": "10-50",
        "competences": ["python", "django"], "certifications": [],
    }, headers=headers)
    payload = {"weights": dict(VALID_WEIGHTS), "threshold_go": 88, "threshold_sous_reserve": 60}
    if business_rules is not None:
        payload["business_rules"] = business_rules
    client.put("/api/scoring-config/policy", json=payload, headers=headers)
    client.post("/api/scoring-config/policy/activate", json={"expected_active_version": None}, headers=headers)
    return user, csrf


def _install_no_op_search(monkeypatch):
    from src.rag import private_rag_manager
    monkeypatch.setattr(private_rag_manager, "search", lambda db, **kw: [])


def _run_analysis_to_terminal(client, csrf, text: str = VALID_AO_TEXT):
    r = client.post("/api/analyze", data={"mode": "paste", "text": text}, headers={"X-CSRF-Token": csrf})
    assert r.status_code == 200, r.text
    job_id = r.json()["job_id"]
    for _ in range(80):
        job = jobs.get_job(job_id)
        if job.status != "running":
            break
        time.sleep(0.1)
    assert job.status != "running", "job never reached a terminal state"
    return job


def _analysis_row(db, job_id: str) -> Analysis:
    db.expire_all()
    return db.execute(select(Analysis).where(Analysis.job_id == job_id)).scalar_one()


def _document_count(db, analysis_id, mime_type=None) -> int:
    stmt = select(func.count()).select_from(AnalysisDocument).where(
        AnalysisDocument.analysis_id == analysis_id
    )
    if mime_type is not None:
        stmt = stmt.where(AnalysisDocument.mime_type == mime_type)
    db.expire_all()
    return db.execute(stmt).scalar_one()


def _resolved_path(storage_path: str):
    from src.web.storage.service import get_storage_service
    return get_storage_service().resolve_for_download(storage_path)


def _pdf_text(path) -> str:
    """Real text out of the real PDF — pypdf is already a test dependency of
    this suite (tests/test_b02_qa_validation.py, tests/test_qa_llm_capture.py
    both use it); no new dependency is introduced."""
    from pypdf import PdfReader
    raw = "".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    return re.sub(r"\s+", " ", raw)


def _docx_text(path) -> str:
    from docx import Document as DocxDocument
    return "\n".join(p.text for p in DocxDocument(str(path)).paragraphs)


class ExplodingLLM:
    """Raises on ANY attribute access — a regeneration that so much as
    *looks* at an LLM adapter fails this test loudly."""

    def __getattr__(self, name):
        raise AssertionError(f"a regeneration must never touch the LLM (attribute accessed: {name!r})")


def _forbid_every_recompute(monkeypatch):
    """After this, any LLM call, any re-scoring and any RAG search raises.
    Installed AFTER the original analysis ran, right before regenerating."""
    from src.agents import llm_client
    from src.agents.scoring_engine import ScoringEngine
    from src.livrables.document_generator import DocumentGenerator
    from src.rag import private_rag_manager

    def _boom(*_a, **_k):
        raise AssertionError("a regeneration must recompute NOTHING — it renders the persisted snapshot only")

    monkeypatch.setattr(llm_client, "ClaudeClient", lambda *a, **k: ExplodingLLM())
    monkeypatch.setattr(DocumentGenerator, "_generate_ai_content", _boom)
    monkeypatch.setattr(ScoringEngine, "score", _boom)
    monkeypatch.setattr(ScoringEngine, "enrich_with_llm", _boom)
    monkeypatch.setattr(private_rag_manager, "search", _boom)


# ---------------------------------------------------------------------------
# Group A — per-document status, reported separately.
# ---------------------------------------------------------------------------

def test_status_reports_pdf_and_docx_separately_and_follows_the_real_files(client, db, monkeypatch):
    _install_no_op_search(monkeypatch)
    user, csrf = _configure_account(client, db, "docstatus@example.com", business_rules=COMPLETE_BUSINESS_RULES)
    org_id = default_org_id(db, user)
    job = _run_analysis_to_terminal(client, csrf)
    assert job.status == "done", job.error

    analysis = _analysis_row(db, job.id)
    status = drs.get_document_status(db, analysis_id=analysis.id, user_id=user.id, organization_id=org_id)
    assert status == {"pdf": "available", "docx": "available"}, (
        "a successful run publishes both documents — both must read available"
    )

    # Simulate the PDF being lost on disk while the analysis row and the
    # document row both survive (the exact case the old all-or-nothing
    # job.files dict could not express).
    pdf_row = db.execute(
        select(AnalysisDocument).where(
            AnalysisDocument.analysis_id == analysis.id, AnalysisDocument.mime_type == drs.PDF_MIME,
        )
    ).scalar_one()
    _resolved_path(pdf_row.storage_path).unlink()

    status = drs.get_document_status(db, analysis_id=analysis.id, user_id=user.id, organization_id=org_id)
    assert status == {"pdf": "unavailable", "docx": "available"}, (
        "a document whose file is gone must never be reported available — the row alone is not proof"
    )
    assert _document_count(db, analysis.id, drs.PDF_MIME) == 1, (
        "the status check must not delete or alter any row"
    )


# ---------------------------------------------------------------------------
# Group B — the ticket's core scenario.
# ---------------------------------------------------------------------------

def test_regeneration_rebuilds_from_the_snapshot_alone_after_a_rendering_failure(client, db, monkeypatch):
    _install_no_op_search(monkeypatch)
    monkeypatch.setattr(
        jobs, "_generate_documents",
        lambda job, ao, result: (_ for _ in ()).throw(RuntimeError("render boom")),
    )
    user, csrf = _configure_account(client, db, "regenafterfail@example.com", business_rules=COMPLETE_BUSINESS_RULES)
    org_id = default_org_id(db, user)
    job = _run_analysis_to_terminal(client, csrf)

    # B11-T1's guarantee, which this ticket builds on.
    assert job.status == "error"
    assert job.error_code == jobs.DOCUMENT_GENERATION_FAILED_ERROR_CODE
    assert job.files == {}
    expected = {
        "score_global": job.result.score_global,
        "decision": job.result.decision,
        "scoring_completeness": job.result.scoring_completeness,
        "scoring_missing": list(job.result.scoring_missing or []),
        "enrichment_status": job.result.enrichment_status,
        "rag_selection_status": job.result.rag_selection_status,
        "data_integrity": job.result.data_integrity,
    }
    expected_policy_version = job.scoring_policy_version

    analysis = _analysis_row(db, job.id)
    assert drs.get_document_status(db, analysis_id=analysis.id, user_id=user.id, organization_id=org_id) == {
        "pdf": "unavailable", "docx": "unavailable",
    }, "nothing was rendered — neither document may be announced as available"

    # Inject markup-shaped text into the PERSISTED snapshot (a controlled
    # fixture: this is text a real AO/LLM could legitimately contain) and
    # keep a business string with an ampersand in it.
    snapshot = dict(analysis.result_data)
    snapshot["ao"] = dict(snapshot["ao"], titre="Refonte SI <b>AT&T</b> & filiales", client="AT&T Services")
    snapshot["result"] = dict(snapshot["result"], forces=["<script>alert(1)</script> reference solide"])
    analysis.result_data = snapshot
    db.commit()

    # From here on, ANY llm/scoring/rag activity is a test failure.
    _forbid_every_recompute(monkeypatch)

    document = drs.regenerate_document(
        db, analysis_id=analysis.id, user_id=user.id, organization_id=org_id, kind="pdf",
    )
    db.commit()

    published = _resolved_path(document.storage_path)
    assert published.exists() and published.stat().st_size > 0
    assert not published.with_name(published.name + ".tmp").exists(), (
        "the atomic publish must leave no temporary artifact behind"
    )
    assert drs.get_document_status(db, analysis_id=analysis.id, user_id=user.id, organization_id=org_id) == {
        "pdf": "available", "docx": "unavailable",
    }, "regenerating the PDF must not claim the DOCX exists"

    # Every scoring field is the ORIGINAL one, field by field — nothing
    # recomputed against the account's current configuration.
    _ao, result = drs._reconstruct_from_snapshot(_analysis_row(db, job.id))
    assert result.score_global == expected["score_global"]
    assert result.decision == expected["decision"]
    assert result.scoring_completeness == expected["scoring_completeness"]
    assert list(result.scoring_missing or []) == expected["scoring_missing"]
    assert result.enrichment_status == expected["enrichment_status"]
    assert result.rag_selection_status == expected["rag_selection_status"]
    assert result.data_integrity == expected["data_integrity"]
    assert _analysis_row(db, job.id).result_data["scoring_policy_version"] == expected_policy_version

    # Markup-shaped text renders as inert literal characters, and real
    # business content with an "&" is fully preserved.
    text = _pdf_text(published)
    assert "<script>alert(1)</script>" in text, "markup must survive as literal, visible text"
    assert "AT&T" in text, "an ampersand in a real client name must not be swallowed by the markup parser"
    assert "<b>AT&T</b>" in text, "deliberate-looking markup in DATA must render literally, never be interpreted"

    # A second regeneration updates the row in place — no duplicate.
    first_path = published
    second = drs.regenerate_document(
        db, analysis_id=analysis.id, user_id=user.id, organization_id=org_id, kind="pdf",
    )
    db.commit()
    assert _document_count(db, analysis.id, drs.PDF_MIME) == 1, (
        "a repeated regeneration must UPDATE the existing AnalysisDocument, never add a second one"
    )
    assert second.id == document.id
    new_path = _resolved_path(second.storage_path)
    assert new_path != first_path, "each publish uses its own fresh target — never a half-overwrite"
    assert new_path.exists()
    assert not first_path.exists(), (
        "the superseded physical file must be deleted once the row points at the new one"
    )


# ---------------------------------------------------------------------------
# Group C — "INCOMPLET" preserved through regeneration.
# ---------------------------------------------------------------------------

def test_incomplet_is_preserved_verbatim_through_regeneration(client, db, monkeypatch):
    """An account with no business_rules configured produces the honest
    "INCOMPLET" verdict (B06-T4). Regenerating its document must render
    that same verdict — never silently resolve it into a real GO/NO-GO."""
    _install_no_op_search(monkeypatch)
    user, csrf = _configure_account(client, db, "incompletregen@example.com")  # business_rules deliberately unset
    org_id = default_org_id(db, user)
    job = _run_analysis_to_terminal(client, csrf)

    assert job.result.decision == "INCOMPLET", (
        "this account has no business rules — the honest verdict is INCOMPLET"
    )
    assert job.result.scoring_missing, "an INCOMPLET verdict must say what is missing"
    expected_missing = list(job.result.scoring_missing)
    expected_score = job.result.score_global

    analysis = _analysis_row(db, job.id)
    _forbid_every_recompute(monkeypatch)

    document = drs.regenerate_document(
        db, analysis_id=analysis.id, user_id=user.id, organization_id=org_id, kind="docx",
    )
    db.commit()

    _ao, result = drs._reconstruct_from_snapshot(_analysis_row(db, job.id))
    assert result.decision == "INCOMPLET", "regeneration must never promote or downgrade the stored decision"
    assert list(result.scoring_missing) == expected_missing
    assert result.score_global == expected_score
    assert result.scoring_completeness == "incomplete"

    text = _docx_text(_resolved_path(document.storage_path))
    assert "INCOMPLET" in text
    assert "Dossier de candidature — WinMarket AI" not in text, (
        "an INCOMPLET analysis must not render as the confident GO dossier"
    )


# ---------------------------------------------------------------------------
# Group D — ownership refused.
# ---------------------------------------------------------------------------

def test_another_user_cannot_regenerate_or_inspect_someone_elses_analysis(client, db, monkeypatch):
    _install_no_op_search(monkeypatch)
    owner, csrf = _configure_account(client, db, "regenowner@example.com", business_rules=COMPLETE_BUSINESS_RULES)
    job = _run_analysis_to_terminal(client, csrf)
    analysis = _analysis_row(db, job.id)
    before = _document_count(db, analysis.id)

    intruder = make_active_starter_user(db, "regenintruder@example.com")
    intruder_org = default_org_id(db, intruder)
    assert intruder.id != owner.id

    with pytest.raises(drs.AnalysisNotFoundError):
        drs.regenerate_document(
            db, analysis_id=analysis.id, user_id=intruder.id, organization_id=intruder_org, kind="pdf",
        )
    with pytest.raises(drs.AnalysisNotFoundError):
        drs.get_document_status(
            db, analysis_id=analysis.id, user_id=intruder.id, organization_id=intruder_org,
        )

    # Exactly the same error type as a genuinely non-existent analysis —
    # the shape itself must not confirm that this analysis exists.
    import uuid as _uuid
    with pytest.raises(drs.AnalysisNotFoundError):
        drs.regenerate_document(
            db, analysis_id=_uuid.uuid4(), user_id=intruder.id, organization_id=intruder_org, kind="pdf",
        )

    assert _document_count(db, analysis.id) == before, "a refused regeneration must write nothing"


def test_a_matching_user_but_mismatched_organization_is_refused(client, db, monkeypatch):
    _install_no_op_search(monkeypatch)
    user, csrf = _configure_account(client, db, "orgmismatch@example.com", business_rules=COMPLETE_BUSINESS_RULES)
    job = _run_analysis_to_terminal(client, csrf)
    analysis = _analysis_row(db, job.id)

    from src.web.database.repositories import organizations as organizations_repo
    other_org = organizations_repo.create_organization(db, name="Autre organisation")
    db.commit()

    with pytest.raises(drs.AnalysisOwnershipMismatchError):
        drs.regenerate_document(
            db, analysis_id=analysis.id, user_id=user.id, organization_id=other_org.id, kind="pdf",
        )


# ---------------------------------------------------------------------------
# Group E — missing / degraded snapshot refused honestly.
# ---------------------------------------------------------------------------

def test_a_snapshot_without_a_result_is_refused_not_filled_in(db):
    from src.web.database.repositories import analyses as analyses_repo

    user = make_active_starter_user(db, "noresultsnap@example.com")
    org_id = default_org_id(db, user)
    analysis = analyses_repo.upsert_analysis(
        db, user_id=user.id, organization_id=org_id, job_id="no-result-job",
        result_data={"id": "no-result-job", "ao": {"titre": "AO sans resultat", "client": "Client"}, "files": {}},
    )
    db.commit()

    with pytest.raises(drs.SnapshotUnusableError):
        drs.regenerate_document(
            db, analysis_id=analysis.id, user_id=user.id, organization_id=org_id, kind="pdf",
        )
    assert _document_count(db, analysis.id) == 0, "no placeholder document may be produced"


def test_a_snapshot_with_a_non_finite_score_is_refused(db):
    """The exact B18-T2 / HISTORICAL_SCORE_UNAVAILABLE case the read path in
    jobs.py already refuses — the regeneration path must refuse it too,
    rather than rendering a document around a NaN."""
    from src.web.database.repositories import analyses as analyses_repo

    user = make_active_starter_user(db, "nanscoresnap@example.com")
    org_id = default_org_id(db, user)
    analysis = analyses_repo.upsert_analysis(
        db, user_id=user.id, organization_id=org_id, job_id="nan-score-job",
        result_data={
            "id": "nan-score-job",
            "ao": {"titre": "AO historique", "client": "Client"},
            "result": {"decision": "GO", "score_global": float("nan"), "criteres": []},
            "files": {},
        },
    )
    db.commit()

    with pytest.raises(drs.SnapshotUnusableError):
        drs.regenerate_document(
            db, analysis_id=analysis.id, user_id=user.id, organization_id=org_id, kind="docx",
        )
    assert _document_count(db, analysis.id) == 0

    # And the stored snapshot is untouched by the refused attempt.
    db.expire_all()
    assert analyses_repo.get_by_job_id(db, "nan-score-job").result_data["result"]["decision"] == "GO"


def test_an_unknown_document_kind_is_refused(db):
    from src.web.database.repositories import analyses as analyses_repo

    user = make_active_starter_user(db, "badkind@example.com")
    org_id = default_org_id(db, user)
    analysis = analyses_repo.upsert_analysis(
        db, user_id=user.id, organization_id=org_id, job_id="bad-kind-job", result_data={},
    )
    db.commit()

    with pytest.raises(drs.UnsupportedDocumentKindError):
        drs.regenerate_document(
            db, analysis_id=analysis.id, user_id=user.id, organization_id=org_id, kind="xlsx",
        )


# ---------------------------------------------------------------------------
# Group F — the escaping fix in document_generator.py.
# ---------------------------------------------------------------------------

def _ao_and_result_with_markup():
    from src.core.models import AOContext, CriterionScore, ScoringResult

    ao = AOContext(titre="Refonte <SI> & data", client="AT&T Services", secteur="Telecom & Media")
    result = ScoringResult(
        decision="NO-GO",
        score_global=42.0,
        criteres=[CriterionScore(
            nom="Rentabilite <estimee>",
            poids=10.0,
            score=30.0,
            justification="budget <50k EUR pour AT&T — marge < 5% & risque eleve",
        )],
        criteres_bloquants=["Budget <50k EUR & hors seuil"],
        forces=["Equipe R&D disponible"],
        faiblesses=["Delai < 3 mois"],
        recommandations=["Renegocier le budget & le delai"],
        ai_content={
            "resume_executif": "Dossier <b>non retenu</b> pour AT&T & filiales.",
            "conclusion": "Marge < 5% : nous ne repondons pas.",
        },
    )
    return ao, result


def test_pdf_renders_markup_characters_as_literal_text_without_losing_content(tmp_path):
    """The defect: ReportLab's Paragraph PARSES a mini-XML markup inside its
    text argument, so an unescaped "&" or "<" from real business data was
    consumed as markup — silently dropping content or breaking the render.
    Every one of these strings must come back out of the finished PDF."""
    from src.livrables.document_generator import DocumentGenerator

    ao, result = _ao_and_result_with_markup()
    target = tmp_path / "out" / "rapport.pdf"
    written = DocumentGenerator(output_dir=tmp_path).generate_pdf(ao, result, llm=None, output_path=target)

    text = _pdf_text(written)
    for expected in [
        "AT&T Services",
        "Refonte <SI> & data",
        "Telecom & Media",
        "Rentabilite <estimee>",
        "budget <50k EUR pour AT&T — marge < 5% & risque eleve",
        "Budget <50k EUR & hors seuil",
        "Equipe R&D disponible",
        "Delai < 3 mois",
        "Renegocier le budget & le delai",
        "Dossier <b>non retenu</b> pour AT&T & filiales.",
        "Marge < 5% : nous ne repondons pas.",
    ]:
        assert expected in text, f"escaping must preserve business content verbatim — missing: {expected!r}"


def test_escape_helper_never_drops_or_truncates_content():
    from src.livrables.document_generator import _escape_for_pdf

    assert _escape_for_pdf("AT&T") == "AT&amp;T"
    assert _escape_for_pdf("<script>alert(1)</script>") == "&lt;script&gt;alert(1)&lt;/script&gt;"
    assert _escape_for_pdf('budget "<50k" & \'marge\'') == "budget &quot;&lt;50k&quot; &amp; &apos;marge&apos;"
    assert _escape_for_pdf(None) == ""
    assert _escape_for_pdf(42) == "42"
    # Nothing but the five markup characters is rewritten, and nothing is lost.
    plain = "Budget 250 000 € — délai 3 mois, équipe 5 ETP (Python/Django)"
    assert _escape_for_pdf(plain) == plain


def test_docx_needs_no_escaping_because_python_docx_stores_text_as_a_literal_node(tmp_path):
    """Confirmed rather than assumed: python-docx assigns the string to an
    lxml TEXT NODE (run._r.text), and lxml escapes a text node on
    serialization — so markup in the value is stored and read back as
    literal characters, never parsed. This is why generate_docx was
    deliberately left untouched by the escaping fix."""
    from docx import Document as DocxDocument
    from src.livrables.document_generator import DocumentGenerator

    # 1) The library-level fact, on its own.
    probe = tmp_path / "probe.docx"
    d = DocxDocument()
    d.add_paragraph("<b>AT&T</b> & <script>alert(1)</script>")
    d.save(str(probe))
    assert "<b>AT&T</b> & <script>alert(1)</script>" in _docx_text(probe), (
        "python-docx must round-trip markup as literal text — if this ever fails, "
        "generate_docx needs the same escaping treatment as generate_pdf"
    )

    # 2) And through the real generator, with the same business data.
    ao, result = _ao_and_result_with_markup()
    written = DocumentGenerator(output_dir=tmp_path).generate_docx(
        ao, result, llm=None, output_path=tmp_path / "out" / "candidature.docx",
    )
    text = _docx_text(written)
    assert "AT&T Services" in text
    assert "budget <50k EUR pour AT&T — marge < 5% & risque eleve" in text
    assert "&amp;" not in text, "the DOCX must contain the real characters, not escaped entities"
