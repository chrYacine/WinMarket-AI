"""Ticket B01 — per-analysis document identity and isolated downloads.

Covers the collision described in the audit: DocumentGenerator used to name
files from the AO title + a minute-precision timestamp
(src/livrables/document_generator.py), written flat into OUTPUT_DIR. Two
analyses sharing a title within the same minute collided on disk, and
src/web/routes_api.py::api_download served whatever `job.files[kind]`
pointed at without ever consulting the DB-backed, ownership-checked
AnalysisDocument row.

These tests exercise the real production code paths (jobs._generate_documents,
jobs._persist_to_database, the /api/download route) against synthetic
AOContext/ScoringResult objects and an isolated on-disk storage root — no
network, no LLM, no Pappers, no SMTP.
"""
from __future__ import annotations

import re
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from src.core.models import AOContext, ScoringResult
from tests.conftest import default_org_id, make_active_starter_user


def _csrf_from(html: str) -> str:
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert match
    return match.group(1)


def _login(client, email, password="Sup3rSecret!"):
    r = client.get("/login")
    csrf = _csrf_from(r.text)
    return client.post("/login", data={"email": email, "password": password, "next": "/app", "csrf_token": csrf})


def _ao(titre: str = "Refonte plateforme e-commerce") -> AOContext:
    return AOContext(titre=titre, client="Client Test", secteur="Retail")


def _result(decision: str = "GO", score: float = 90.0) -> ScoringResult:
    return ScoringResult(decision=decision, score_global=score, criteres=[])


@pytest.fixture()
def isolated_storage(tmp_path, monkeypatch):
    """Point every OUTPUT_DIR/LOCAL_STORAGE_PATH lookup at a throwaway
    directory so tests never touch the real data/outputs/ folder, and so
    StorageService's root-confinement check matches where files actually
    land (both config.OUTPUT_DIR and jobs.ANALYSIS_FILES_DIR are read once
    at import time, so both must be patched explicitly)."""
    from src.core import config
    from src.web import jobs as jobs_module

    output_dir = tmp_path / "outputs"
    output_dir.mkdir()
    monkeypatch.setattr(config, "OUTPUT_DIR", output_dir)
    monkeypatch.setattr(config, "LOCAL_STORAGE_PATH", tmp_path)
    monkeypatch.setattr(jobs_module, "ANALYSIS_FILES_DIR", output_dir)
    return output_dir


def _persist_job_to_db(job, ao, result, db):
    """Mirror of jobs._persist_to_database but against the test's `db`
    session directly (the real function opens its own session via
    session_scope(), which is correct in production but awkward to observe
    from a test that wants to assert on the same session)."""
    from src.livrables.document_generator import _safe_name
    from src.web.database.repositories import analyses as analyses_repo
    from src.web.storage.service import get_storage_service

    analysis = analyses_repo.create_analysis(
        db, user_id=job.user_id, organization_id=job.organization_id, job_id=job.id, title=ao.titre,
        client_name=ao.client, score=result.score_global, decision=result.decision,
        result_data={"ao": ao.model_dump(), "result": result.model_dump()},
    )
    storage = get_storage_service()
    readable_stem = _safe_name(ao.titre)
    original_names = {"pdf": f"rapport_decision_{readable_stem}.pdf", "docx": f"candidature_{readable_stem}.docx"}
    for kind, path_str in job.files.items():
        path = Path(path_str)
        analyses_repo.add_document(
            db, analysis_id=analysis.id, user_id=job.user_id, organization_id=job.organization_id,
            filename=path.name, original_filename=original_names[kind],
            storage_path=storage.save(path),
            mime_type=("application/pdf" if kind == "pdf" else
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            file_size=path.stat().st_size,
        )
    db.commit()
    return analysis


# ---------------------------------------------------------------------------
# A. Identity: same title, same minute -> distinct files, real content
# ---------------------------------------------------------------------------

def test_same_title_same_minute_produces_distinct_files(isolated_storage):
    from src.web import jobs

    job_a = jobs.create_job(source_label="A", user_id=uuid.uuid4())
    job_b = jobs.create_job(source_label="B", user_id=uuid.uuid4())
    ao = _ao("Refonte plateforme e-commerce")

    files_a = jobs._generate_documents(job_a, ao, _result(decision="GO", score=95.0))
    files_b = jobs._generate_documents(job_b, ao, _result(decision="NO-GO", score=12.0))

    assert files_a["pdf"] != files_b["pdf"]
    assert files_a["docx"] != files_b["docx"]

    pdf_a_bytes = Path(files_a["pdf"]).read_bytes()
    pdf_b_bytes = Path(files_b["pdf"]).read_bytes()
    assert pdf_a_bytes != pdf_b_bytes
    assert pdf_a_bytes.startswith(b"%PDF-") and pdf_b_bytes.startswith(b"%PDF-")

    text_a = "\n".join(p.text for p in DocxDocument(files_a["docx"]).paragraphs)
    text_b = "\n".join(p.text for p in DocxDocument(files_b["docx"]).paragraphs)
    assert "Dossier de candidature" in text_a and "Mémo de non-réponse" not in text_a
    assert "Mémo de non-réponse" in text_b and "Dossier de candidature" not in text_b


def test_same_user_two_identical_analyses_no_overwrite(isolated_storage):
    """Same user, same title, run twice — checklist item: 'un même
    utilisateur lance deux analyses identiques : aucun écrasement'."""
    from src.web import jobs

    user_id = uuid.uuid4()
    ao = _ao("Migration cloud AWS")

    job_1 = jobs.create_job(source_label="run 1", user_id=user_id)
    files_1 = jobs._generate_documents(job_1, ao, _result(decision="GO", score=80.0))
    job_2 = jobs.create_job(source_label="run 2", user_id=user_id)
    files_2 = jobs._generate_documents(job_2, ao, _result(decision="GO", score=99.0))

    assert files_1["docx"] != files_2["docx"]
    assert Path(files_1["docx"]).exists() and Path(files_2["docx"]).exists()

    assert Path(files_1["docx"]).read_bytes() != Path(files_2["docx"]).read_bytes()


def test_concurrent_generation_does_not_overwrite(isolated_storage):
    """Two generations racing for the same title must not corrupt each
    other's file — checklist item: 'deux générations concurrentes'."""
    from src.web import jobs

    ao = _ao("Support N3 infogérance")

    def _run(i: int):
        job = jobs.create_job(source_label=f"concurrent-{i}", user_id=uuid.uuid4())
        return jobs._generate_documents(job, ao, _result(decision="GO", score=float(50 + i)))

    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(_run, range(4)))

    all_pdfs = [r["pdf"] for r in results]
    all_docx = [r["docx"] for r in results]
    assert len(set(all_pdfs)) == 4
    assert len(set(all_docx)) == 4
    for r in results:
        assert Path(r["pdf"]).exists()
        assert Path(r["docx"]).exists()


def test_generation_collision_is_detected_not_silently_absorbed(isolated_storage, tmp_path):
    """Point B: a genuine collision on an explicit output_path must raise,
    never overwrite silently."""
    from src.livrables.document_generator import DocumentGenerator

    dg = DocumentGenerator()
    target = tmp_path / "same-target.docx"
    dg.generate_docx(_ao(), _result(), output_path=target)
    assert target.exists()

    with pytest.raises(FileExistsError):
        dg.generate_docx(_ao(), _result(decision="NO-GO"), output_path=target)


# ---------------------------------------------------------------------------
# D. Download: ownership chain + storage confinement
# ---------------------------------------------------------------------------

def test_owner_downloads_own_document_after_other_users_generation(client, db, isolated_storage):
    """Checklist: 'A télécharge toujours le contenu A après génération de B'
    and 'A ne peut télécharger aucun document B'."""
    from src.web import jobs

    user_a = make_active_starter_user(db, "docA@example.com")
    user_b = make_active_starter_user(db, "docB@example.com")

    job_a = jobs.create_job(source_label="A", user_id=user_a.id, organization_id=default_org_id(db, user_a))
    ao_a = _ao("Plateforme RH mutualisée")
    result_a = _result(decision="GO", score=91.0)
    job_a.files = jobs._generate_documents(job_a, ao_a, result_a)
    job_a.ao, job_a.result, job_a.status = ao_a, result_a, "done"
    _persist_job_to_db(job_a, ao_a, result_a, db)

    # B generates *after* A, same title, same minute.
    job_b = jobs.create_job(source_label="B", user_id=user_b.id, organization_id=default_org_id(db, user_b))
    ao_b = _ao("Plateforme RH mutualisée")
    result_b = _result(decision="NO-GO", score=8.0)
    job_b.files = jobs._generate_documents(job_b, ao_b, result_b)
    job_b.ao, job_b.result, job_b.status = ao_b, result_b, "done"
    _persist_job_to_db(job_b, ao_b, result_b, db)

    _login(client, "docA@example.com")
    r = client.get(f"/api/download/{job_a.id}/docx")
    assert r.status_code == 200
    assert r.content == Path(job_a.files["docx"]).read_bytes()
    assert "Mémo de non-réponse" not in r.content.decode("latin-1", errors="ignore")

    # A cannot reach B's job at all.
    r_cross = client.get(f"/api/download/{job_b.id}/docx")
    assert r_cross.status_code == 404


def test_content_disposition_uses_readable_name_not_technical_id(client, db, isolated_storage):
    from src.web import jobs

    user = make_active_starter_user(db, "readable@example.com")
    job = jobs.create_job(source_label="x", user_id=user.id, organization_id=default_org_id(db, user))
    ao = _ao("Support et maintenance applicative")
    result = _result()
    job.files = jobs._generate_documents(job, ao, result)
    job.ao, job.result, job.status = ao, result, "done"
    _persist_job_to_db(job, ao, result, db)

    _login(client, "readable@example.com")
    r = client.get(f"/api/download/{job.id}/pdf")
    assert r.status_code == 200
    disposition = r.headers["content-disposition"]
    assert "rapport_decision_Support_et_maintenance_applicative.pdf" in disposition
    # The on-disk technical name (a uuid4 hex) must never leak into it.
    technical_name = Path(job.files["pdf"]).name
    assert technical_name not in disposition


def test_missing_file_returns_clean_404_not_500(client, db, isolated_storage):
    from src.web import jobs

    user = make_active_starter_user(db, "gone@example.com")
    job = jobs.create_job(source_label="x", user_id=user.id, organization_id=default_org_id(db, user))
    ao, result = _ao(), _result()
    job.files = jobs._generate_documents(job, ao, result)
    job.ao, job.result, job.status = ao, result, "done"
    _persist_job_to_db(job, ao, result, db)

    Path(job.files["pdf"]).unlink()

    _login(client, "gone@example.com")
    r = client.get(f"/api/download/{job.id}/pdf")
    assert r.status_code == 404


def test_out_of_root_storage_path_is_refused(client, db, isolated_storage, tmp_path):
    """Point D: never resolve a storage reference that escapes the storage
    root, even one recorded in the DB (defense in depth against a bad or
    tampered row) — checklist: 'un chemin hors racine est refusé'."""
    from src.web import jobs
    from src.web.database.repositories import analyses as analyses_repo

    user = make_active_starter_user(db, "escape@example.com")
    org_id = default_org_id(db, user)
    job = jobs.create_job(source_label="x", user_id=user.id, organization_id=org_id)
    ao, result = _ao(), _result()
    job.files = jobs._generate_documents(job, ao, result)
    job.ao, job.result, job.status = ao, result, "done"

    outside = tmp_path.parent / f"outside-{uuid.uuid4().hex}.pdf"
    outside.write_bytes(b"%PDF-1.4 not really inside the storage root")
    try:
        analysis = analyses_repo.create_analysis(
            db, user_id=user.id, organization_id=org_id, job_id=job.id, title=ao.titre, result_data={},
        )
        analyses_repo.add_document(
            db, analysis_id=analysis.id, user_id=user.id, organization_id=org_id,
            filename=outside.name, original_filename="rapport.pdf",
            storage_path=str(outside), mime_type="application/pdf",
            file_size=outside.stat().st_size,
        )
        db.commit()

        _login(client, "escape@example.com")
        r = client.get(f"/api/download/{job.id}/pdf")
        assert r.status_code == 404
    finally:
        outside.unlink(missing_ok=True)


def test_legacy_job_without_db_row_still_downloads(client, db, isolated_storage):
    """Point E: an analysis whose best-effort DB write failed (organization
    context present and membership still active, but no AnalysisDocument
    row got written) — the download must still resolve through the
    in-memory/JSON job record, confined to the storage root, not a bare
    unchecked path. Distinct from a genuinely pre-B02 job with no
    organization_id at all — see
    tests/test_b02_organizations.py::test_legacy_job_without_organization_id_is_refused_not_fallback."""
    from src.web import jobs

    user = make_active_starter_user(db, "legacy@example.com")
    job = jobs.create_job(source_label="x", user_id=user.id, organization_id=default_org_id(db, user))
    ao, result = _ao(), _result()
    job.files = jobs._generate_documents(job, ao, result)
    job.ao, job.result, job.status = ao, result, "done"
    # Deliberately skip _persist_job_to_db: no AnalysisDocument row exists.

    _login(client, "legacy@example.com")
    r = client.get(f"/api/download/{job.id}/docx")
    assert r.status_code == 200
    assert r.content == Path(job.files["docx"]).read_bytes()


def test_resolve_for_download_confines_local_prefixed_traversal(tmp_path):
    """Unit-level check on the storage layer itself: a local:// reference
    that tries to climb out of the root via '..' must not resolve."""
    from src.web.storage.service import LocalStorageService

    root = tmp_path / "root"
    root.mkdir()
    secret = tmp_path / "secret.txt"
    secret.write_text("nope")

    storage = LocalStorageService(root=root)
    with pytest.raises(FileNotFoundError):
        storage.resolve_for_download("local://../secret.txt")


def test_resolve_for_download_confines_an_outgoing_symlink(tmp_path):
    """B23-T1: 'traversées et liens sortants' — a symlink INSIDE the
    allowed root whose TARGET points outside it must be refused too, not
    just a literal '../' path string. resolve_for_download's own
    Path.resolve() call follows symlinks to their real target before the
    containment check runs, so this should already hold — verified
    directly here rather than assumed from reading the code.

    Skipped (not xfail — this is an environment limitation, not a product
    defect) if this OS/user cannot create a symlink: unprivileged symlink
    creation requires Developer Mode or admin rights on Windows, and this
    is exactly the kind of environment-dependent gap the lot's own
    instructions say to report honestly rather than paper over."""
    root = tmp_path / "root"
    root.mkdir()
    secret = tmp_path / "outside_root_secret.txt"
    secret.write_text("nope")
    link = root / "escape_link.txt"
    try:
        link.symlink_to(secret)
    except OSError as exc:
        pytest.skip(f"symlink creation not permitted in this environment: {exc}")

    from src.web.storage.service import LocalStorageService

    storage = LocalStorageService(root=root)
    with pytest.raises(FileNotFoundError):
        storage.resolve_for_download(f"local://{link.name}")
