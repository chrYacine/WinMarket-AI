import os
import re
from pathlib import Path
from datetime import datetime
from typing import Optional
from xml.sax.saxutils import escape as _xml_escape
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import cm
from src.agents.certification_scope import _NAME_PATTERNS
from src.core.config import OUTPUT_DIR, LLM_TEMPERATURE_GENERATION
from src.core.logger import get_agent_logger
from src.core.models import AOContext, ScoringResult
from src.core.prompt_loader import load_prompt

_logger = get_agent_logger("document_generator")

# B19-T1: the four active generation prompts live in their own text files
# and are loaded through the ONE shared loader (src/core/prompt_loader.py),
# exactly as src/rag/prompts/reference_selection.txt (B18-T3) and
# src/agents/prompts/ao_extraction_*.txt (B05-T2) already do — never a
# second ad hoc loader, never an inline f-string prompt.
_PROMPTS_DIR = Path(__file__).parent / "prompts"
_SYSTEM_PATH = _PROMPTS_DIR / "document_system.txt"
_PROMPT_PATHS = {
    "NO-GO": _PROMPTS_DIR / "document_no_go.txt",
    "RESERVE": _PROMPTS_DIR / "document_reserve.txt",
    "GO": _PROMPTS_DIR / "document_go.txt",
}

# B19-T1: "PRIS" is a genuine ANSSI qualification AND the single most
# common French past participle ("nous avons pris en compte..."), so
# _NAME_PATTERNS' case-insensitive `\bpris\b` matches ordinary prose. This
# does NOT redefine the certification vocabulary (which stays the single
# copy in src/agents/certification_scope.py) — it only disambiguates a
# match of these names: a lowercase occurrence is the common word, an
# uppercase one is the certification.
_CASE_SENSITIVE_NAMES = {"PRIS"}


def _certifications_mentioned(text: str) -> set[str]:
    """Canonical certification names mentioned in `text`, using the SAME
    vocabulary the extraction pipeline recognizes."""
    found: set[str] = set()
    for name, pattern in _NAME_PATTERNS.items():
        for match in re.finditer(pattern, text, re.IGNORECASE):
            if name in _CASE_SENSITIVE_NAMES and not match.group(0).isupper():
                continue
            found.add(name)
            break
    return found


def _known_certifications(ao: AOContext, provider_profile=None) -> set[str]:
    """Lowercased set of certification names the SERVER actually knows to
    be real for this analysis: the ones the AO itself requires, plus the
    ones this account declared on its own ProviderProfile. Each raw string
    is kept as-is AND canonicalized through the shared vocabulary, so a
    free-typed "iso27001" matches a generated "ISO 27001"."""
    raw: list[str] = [str(c) for c in (ao.certifications_obligatoires or [])]
    if provider_profile is not None:
        for entry in (getattr(provider_profile, "certifications", None) or []):
            nom = str(entry.get("nom", "")).strip() if isinstance(entry, dict) else str(entry).strip()
            if nom:
                raw.append(nom)

    known: set[str] = set()
    for item in raw:
        item = item.strip()
        if not item:
            continue
        known.add(item.lower())
        known.update(n.lower() for n in _certifications_mentioned(item))
    return known


def _safe_name(titre: str) -> str:
    return "".join(c if c.isalnum() else "_" for c in titre[:40]) or "ao"


# B19-T2 (DEFECT confirmed): ReportLab's `Paragraph` does not take plain
# text — it PARSES a mini-XML/HTML-like markup inside its text argument
# (<b>, <i>, <br/>, <font ...>, <super>, ...). generate_pdf below built
# those Paragraphs by f-string-interpolating free-form values straight from
# the AO, the scoring result and the LLM-generated ai_content, so a literal
# "<" in real business data ("budget <50k€") or a stray "&" ("AT&T") was
# handed to that parser as markup: at best the characters vanished from the
# rendered document, at worst the whole render raised.
#
# Fix: escape every DYNAMIC value at the point it enters a Paragraph, with
# stdlib xml.sax.saxutils.escape (no new dependency — reportlab.lib.utils
# has no general-purpose escape helper of its own in this version; its
# paraparser is an XML parser, so the stdlib XML escaper is exactly the
# right tool). The extra '"'/"'" entities are included because ReportLab's
# markup also treats quotes specially inside tag ATTRIBUTES — escaping them
# costs nothing for ordinary prose and closes that case too.
#
# This is purely additive and strictly content-PRESERVING: only the five
# characters a markup parser would otherwise consume are rewritten into
# their entity form, which ReportLab renders back as the original literal
# character. No text is ever dropped or truncated here. Literal strings
# written in this file (section titles, the deliberate "<b>...</b>" wrapper
# around a criterion name) are NOT passed through it — they are markup we
# actually mean.
#
# generate_docx deliberately gets NO equivalent: python-docx's
# add_paragraph(text)/add_run(text) assign to an lxml text node
# (`run._r.text`), and lxml serializes a text node by escaping it — markup
# in the value is stored and re-read as literal characters, never parsed.
# See tests/test_b19_t2_document_rendering.py, which asserts that
# round-trip rather than assuming it.
_PDF_QUOTE_ESCAPES = {'"': "&quot;", "'": "&apos;"}


def _escape_for_pdf(value) -> str:
    """Make an arbitrary dynamic value safe to interpolate into a ReportLab
    `Paragraph`, without losing a single character of real content."""
    if value is None:
        return ""
    return _xml_escape(value if isinstance(value, str) else str(value), _PDF_QUOTE_ESCAPES)


def _prepare_target(path: Path) -> Path:
    """Resolve the final write target for a generated document.

    Callers that pass an explicit `output_path` (the web app — see
    src/web/jobs.py) get a unique path per document (job/user/document id),
    so a collision here means two callers were handed the *same* target by
    mistake — that must fail loudly, never overwrite silently.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise FileExistsError(f"Un document existe déjà à cet emplacement : {path}")
    return path


def _finalize(tmp_path: Path, final_path: Path) -> Path:
    """Atomically publish a fully-written temp file as `final_path`, so a
    reader can never observe a partially-written document."""
    os.replace(tmp_path, final_path)
    return final_path


def _criterion_state(c) -> str:
    """Lot 44: `etat` is None on a result stored before it (displayed as a
    plain evaluated criterion, exactly as before)."""
    return getattr(c, "etat", None) or "evalue"


def _criterion_text(c) -> str:
    """One criterion as a line of a document: evaluated / hypothesis / not
    evaluated / not applicable — never a fabricated score."""
    state = _criterion_state(c)
    if state == "manquant":
        return f"NON ÉVALUÉ — {c.justification}"
    if state == "non_applicable":
        return f"NON APPLICABLE — {c.justification}"
    hypo = " (hypothèse de la politique, non établie par l'AO)" if state == "hypothese" else ""
    return f"{c.score:.0f}/100{hypo} — {c.justification}"


def _score_line(result) -> str:
    return f"{result.score_global}/100" + (" — provisoire, des critères ne sont pas évalués" if getattr(result, "score_provisoire", None) else "")


_DOSSIER_FIELD_LABELS = {
    "budget_estime": "Budget estimé", "duree_projet_mois": "Durée du projet (mois)", "deadline_reponse": "Date limite de réponse",
    "titre": "Titre", "client": "Client", "secteur": "Secteur", "certifications_obligatoires": "Certification obligatoire",
}


def _dossier_source(src) -> str:
    """One provenance reference: category, safe display name, page when the format has pages, analysed passage."""
    text = f"{src.get('categorie_libelle', '')} — {src.get('nom', '')}"
    passage = src.get("passage") or {}
    pages = passage.get("pages")
    if pages:
        text += f", page {pages[0]}" if pages[0] == pages[1] else f", pages {pages[0]}–{pages[1]}"
    if passage:
        text += f" (passage {passage.get('fenetre')}/{passage.get('sur')})"
    return text


def _shown_value(champ, value, unit=None) -> str:
    """A value as a reader wants it: a budget in euros, a list joined, an integer without ".0"."""
    if champ == "budget_estime" and isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{value:,.0f} €".replace(",", " ")
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return f"{value}" + (f" {unit}" if unit else "")


_CATEGORY_LABELS_FR = {"rc": "RC – Règlement de consultation", "cctp": "CCTP / Cahier des charges",
                       "ccap": "CCAP / Conditions contractuelles", "acte_engagement": "Acte d'engagement"}


def _dossier_scope_note(dossier: dict) -> str | None:
    """Lot 50 §4/§5 — the SAME distinct scope indication as the result page's own banner, carried into the
    deliverables: never merged with the INCOMPLET/decision wording, never claiming an exhaustive validation."""
    if not dossier.get("perimetre_limite"):
        return None
    missing = dossier.get("categories_manquantes") or []
    admitted = sum(1 for p in dossier.get("pieces", []) if p.get("sera_pris_en_compte", True))
    total = len(dossier.get("pieces", []))
    note = f"Analyse limitée aux pièces fournies ({admitted}/{total} pièce(s) analysée(s))."
    if missing:
        note += " Catégorie(s) non fournie(s) : " + ", ".join(_CATEGORY_LABELS_FR.get(c, c) for c in missing) + "."
    note += " Ce résultat ne constitue pas une validation exhaustive du marché."
    return note


def _dossier_lines(ao) -> tuple[list[str], list[str], list[str], str | None]:
    """(pieces, conflicts, provenance, scope_note) of a consolidated AO dossier as plain lines — empty/None
    for a single-file analysis. Only the public description is used: never an internal path or storage key."""
    dossier = getattr(ao, "dossier", None)
    if not dossier or not isinstance(dossier, dict):
        return [], [], [], None
    pieces = []
    for p in dossier.get("pieces", []):
        line = f"{p.get('categorie_libelle', '')} — {p.get('nom', '')} ({str(p.get('format', '')).upper()}, {p.get('taille', '')})"
        if p.get("doublon_de"):
            line += " — doublon, non compté deux fois"
        if p.get("sera_pris_en_compte") is False:
            line += " — EXCLUE" + (f" ({p['raison_exclusion']})" if p.get("raison_exclusion") else "")
        pieces.append(line)
    conflicts = []
    for c in dossier.get("conflits", []):
        values = "; ".join(
            _shown_value(c["champ"], v.get("valeur"), v.get("unite"))
            + (" [" + ", ".join(_dossier_source(s) for s in v["sources"]) + "]" if v.get("sources") else "")
            for v in c.get("valeurs", [])
        )
        conflicts.append(f"{c.get('libelle') or _DOSSIER_FIELD_LABELS.get(c['champ'], c['champ'])} — {c.get('resolution', '')} : {values}")
    provenance = []
    for o in dossier.get("observations", []):
        if o["champ"] in ("titre", "client", "secteur"):  # display data, never a scoring input
            continue
        srcs = [_dossier_source(o["source"])] + [_dossier_source(s) for s in o.get("autres_sources", [])]
        provenance.append(f"{o.get('libelle') or _DOSSIER_FIELD_LABELS.get(o['champ'], o['champ'])} : {_shown_value(o['champ'], o.get('valeur'), o.get('unite'))} — {' ; '.join(srcs)}")
    return pieces, conflicts, provenance, _dossier_scope_note(dossier)


_ORIGIN_LABEL_PLAIN = {"declared_user": "déclaré par vous", "llm_sourced": "proposition acceptée, citation retrouvée"}


def _shown_change_value(value) -> str:
    if value is None:
        return "Manquante"
    if isinstance(value, list):
        return ", ".join(str(v) for v in value) if value else "Manquante"
    if isinstance(value, bool):
        return "oui" if value else "non"
    return str(value)


def _completion_change_lines(result: ScoringResult) -> list[str]:
    """Lot 53 — the SAME `completion_changes` the result page's "Ce qui a changé" table renders (frozen at
    submission time by src/web/completion_service.py, never recomputed here) as plain text lines, shared by
    both PDF and DOCX. Empty for an ordinary (non-revision) analysis, and for any revision computed before
    this field existed (never fabricated after the fact)."""
    lines = []
    for ch in getattr(result, "completion_changes", None) or []:
        origin = _ORIGIN_LABEL_PLAIN.get(ch.get("origin"), ch.get("origin") or "")
        source = ch.get("source_json") if isinstance(ch.get("source_json"), dict) else None
        citation = f" — citation : « {source['citation']} »" if source and source.get("citation") else ""
        lines.append(f"{ch.get('field', '?')} : {_shown_change_value(ch.get('before'))} → {_shown_change_value(ch.get('after'))} ({origin}{citation})")
    return lines


class DocumentGenerator:
    def __init__(self, output_dir: Path = OUTPUT_DIR):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _build_doc_system_prompt(self, provider_profile=None) -> str:
        """B19-T1 (DEFECT confirmed): the system prompt used to hardcode a
        FICTIONAL company identity — "une ESN française spécialisée en
        transformation digitale, développement applicatif, data/IA, cloud
        et cybersécurité" — a specialization no account ever declared. The
        identity sentence is now built from the account's OWN
        ProviderProfile (real raison_sociale, real declared competences)
        when one is given, and is otherwise deliberately specific-free: no
        headcount, no track record, no specialization is ever asserted on
        behalf of a provider that did not declare it. The rest of the
        prompt (tone/quality guidance, which states no fact about any
        specific company) is unchanged, in
        src/livrables/prompts/document_system.txt."""
        raison = ""
        if provider_profile is not None:
            raison = (getattr(provider_profile, "raison_sociale", None) or "").strip()

        if raison:
            competences = [
                str(c).strip()
                for c in (getattr(provider_profile, "competences", None) or [])
                if str(c).strip()
            ]
            if competences:
                identity = (
                    f"Tu rédiges la réponse de {raison} à cet appel d'offres. Les compétences que {raison} a "
                    f"elle-même déclarées sont : {', '.join(competences[:15])}. N'attribue à {raison} "
                    f"aucune autre compétence, aucune taille, aucune ancienneté et aucune réalisation "
                    f"passée qui ne te soit explicitement communiquée."
                )
            else:
                identity = (
                    f"Tu rédiges la réponse de {raison} à cet appel d'offres. Aucune compétence, taille, "
                    f"ancienneté ni réalisation passée de {raison} ne t'est communiquée : n'en invente "
                    f"aucune."
                )
        else:
            identity = (
                "Tu rédiges la réponse du prestataire à cet appel d'offres. "
                "Son identité, sa taille, ses spécialités et ses réalisations passées ne te sont pas "
                "communiquées : n'en invente aucune et n'attribue au prestataire que ce qui figure "
                "explicitement dans le contexte fourni."
            )
        return load_prompt(_SYSTEM_PATH, identity=identity)

    def _build_context(self, ao: AOContext, result: ScoringResult) -> str:
        refs_text = "\n".join([f"  • {ev.source} — pertinence {min(int(ev.score*400),100)}%" for ev in result.evidence_pack[:4]]) or "  • Aucune référence disponible"
        return f"""APPEL D'OFFRES
Titre     : {ao.titre}
Client    : {ao.client}
Secteur   : {ao.secteur or 'Non précisé'}
Budget    : {f"{ao.budget_estime:,.0f} €" if ao.budget_estime else 'Non communiqué'}
Durée     : {f"{ao.duree_projet_mois} mois" if ao.duree_projet_mois else 'Non précisée'}
Techs     : {', '.join(ao.technologies_demandees) or 'Non précisées'}
Livrables : {', '.join(ao.livrables[:4]) or 'Non précisés'}
Contraintes: {' | '.join(ao.contraintes[:3]) or 'Aucune'}
Certif.   : {', '.join(ao.certifications_obligatoires) or 'Aucune'}
Deadline  : {ao.deadline_reponse or 'Non précisée'}

SCORING
Décision  : {result.decision} ({_score_line(result)})
Bloquants : {' | '.join(result.criteres_bloquants) or 'Aucun'}
Forces    : {' | '.join(result.forces[:3]) or 'Aucune'}
Faiblesses: {' | '.join(result.faiblesses[:3]) or 'Aucune'}
Références mobilisables :
{refs_text}"""

    def _build_prompt(self, ao: AOContext, result: ScoringResult) -> str:
        """The decision-specific prompt body, loaded from its own text file.
        `_build_context` stays Python — it assembles DATA, not instructions;
        only the instructional text itself lives in the .txt files."""
        ctx = self._build_context(ao, result)
        decision = result.decision

        # Integration fix (B06-T4 introduces "INCOMPLET" — a scoring result
        # missing a required business rule, never a confident verdict):
        # anything that isn't clearly "GO" or a "RESERVE" variant renders
        # through the NO-GO-shaped, most conservative path — never the
        # confident GO dossier that used to be this branch's silent
        # catch-all. A real NO-GO still matches here exactly as before.
        if decision != "GO" and "RESERVE" not in decision:
            return load_prompt(_PROMPT_PATHS["NO-GO"], context=ctx, client=ao.client)
        if "RESERVE" in decision:
            return load_prompt(
                _PROMPT_PATHS["RESERVE"],
                context=ctx,
                client=ao.client,
                secteur=ao.secteur or "",
                deadline=ao.deadline_reponse or "la date limite",
                livrables=", ".join(ao.livrables[:3]) or "non précisés",
                technologies=", ".join(ao.technologies_demandees[:3]) or "demandée",
                secteur_equipe=ao.secteur or "du client",
            )
        return load_prompt(
            _PROMPT_PATHS["GO"],
            context=ctx,
            client=ao.client,
            secteur=ao.secteur or "",
            technologies=", ".join(ao.technologies_demandees[:3]) or "stack demandée",
            livrables=", ".join(ao.livrables[:3]) or "non précisés",
        )

    def _drop_ungrounded_fields(self, data: dict, ao: AOContext, provider_profile=None) -> dict:
        """B19-T1: consistency guard, not just shape validation. A generated
        field that names a certification the server does NOT know to be real
        for this analysis is dropped WHOLE, so the caller's already-existing
        `ai.get(field) or <neutral fallback>` pattern in generate_docx /
        generate_pdf supplies its own grounded sentence instead.

        "Known real" = the certifications the AO itself requires
        (ao.certifications_obligatoires) UNION the ones this account
        declared on its own ProviderProfile. Deliberately NOT a general
        semantic fact-checker: certification NAMES are the one class of
        claim this codebase can verify cheaply and precisely, using the
        same vocabulary the extraction pipeline uses.

        Logged server-side only — a dropped field is never signalled inside
        the generated document."""
        known = _known_certifications(ao, provider_profile)
        cleaned = dict(data)
        for field, value in data.items():
            if isinstance(value, str):
                texts = [value]
            elif isinstance(value, list):
                texts = [v for v in value if isinstance(v, str)]
            else:
                continue
            mentioned: set[str] = set()
            for text in texts:
                mentioned |= _certifications_mentioned(text)
            invented = sorted(n for n in mentioned if n.lower() not in known)
            if invented:
                cleaned.pop(field, None)
                _logger.warning(
                    "Champ '%s' écarté du contenu généré : certification(s) non justifiée(s) par "
                    "les données du serveur — %s (certifications connues : %s)",
                    field, ", ".join(invented), ", ".join(sorted(known)) or "aucune",
                )
        return cleaned

    def _generate_ai_content(self, ao: AOContext, result: ScoringResult, llm, provider_profile=None) -> dict:
        """Génère le contenu IA adapté à la décision (GO / GO SOUS RÉSERVE / NO-GO).

        `provider_profile` (a src.web.database.models.ProviderProfile row, or
        None) is the account's OWN declared identity."""
        if not llm.enabled:
            return {}

        prompt = self._build_prompt(ao, result)
        system = self._build_doc_system_prompt(provider_profile)

        try:
            data = llm.json_complete(prompt, system=system, temperature=LLM_TEMPERATURE_GENERATION, max_tokens=3000) or {}
        except Exception:
            return {}
        if not isinstance(data, dict):
            return {}
        return self._drop_ungrounded_fields(data, ao, provider_profile)

    def generate_docx(self, ao: AOContext, result: ScoringResult, llm=None, *, output_path: Optional[Path] = None) -> Path:
        if output_path is None:
            # Legacy flat naming — kept exactly as-is for callers that don't
            # scope their own path.
            safe = _safe_name(ao.titre)
            ts = datetime.now().strftime("%Y%m%d_%H%M")
            path = self.output_dir / f"candidature_{safe}_{ts}.docx"
        else:
            path = _prepare_target(Path(output_path))
        tmp_path = path.with_name(path.name + ".tmp")

        ai = result.ai_content or (self._generate_ai_content(ao, result, llm) if llm else {})
        decision = result.decision

        doc = Document()

        # ── Header commun ──
        # Same conservative-default fix as _build_prompt above: only an
        # exact "GO" gets the confident dossier heading.
        if decision != "GO" and "RESERVE" not in decision:
            doc.add_heading("Mémo de non-réponse — WinMarket AI", 0)
        elif "RESERVE" in decision:
            doc.add_heading("Dossier de candidature conditionnel — WinMarket AI", 0)
        else:
            doc.add_heading("Dossier de candidature — WinMarket AI", 0)

        p = doc.add_paragraph()
        p.add_run("Décision : ").bold = True
        p.add_run(result.decision + "   ")
        p.add_run("Score : ").bold = True
        p.add_run(f"{result.score_global}/100   |   ")
        p.add_run("Client : ").bold = True
        p.add_run(f"{ao.client}   |   {ao.titre[:80]}")

        # Lot 47 bis — a dossier analysis says which pieces it read and where the facts come from.
        d_pieces, d_conflicts, d_provenance, d_scope_note = _dossier_lines(ao)
        if d_pieces:
            doc.add_heading("Dossier d'appel d'offres analysé", 1)
            if d_scope_note:  # lot 50 §4/§5 — distinct from the decision text above, never merged with it
                doc.add_paragraph().add_run(d_scope_note).italic = True
            for line in d_pieces:
                doc.add_paragraph(line, style="List Bullet")
            if d_conflicts:
                doc.add_paragraph().add_run("Valeurs contradictoires entre les pièces (non tranchées) :").bold = True
                for line in d_conflicts:
                    doc.add_paragraph(line, style="List Bullet")
            if d_provenance:
                doc.add_paragraph().add_run("Provenance des informations relevées :").bold = True
                for line in d_provenance:
                    doc.add_paragraph(line, style="List Bullet")

        # Lot 53 — this revision's own before/after (frozen at submission time, never recomputed here).
        change_lines = _completion_change_lines(result)
        if change_lines:
            doc.add_heading("Compléments et changements de cette révision", 1)
            for line in change_lines:
                doc.add_paragraph(line, style="List Bullet")

        # ════════════════════════════════
        # CAS NO-GO (et tout ce qui n'est ni GO ni RESERVE, y compris
        # "INCOMPLET" — voir le commentaire sur le header ci-dessus)
        # ════════════════════════════════
        if decision != "GO" and "RESERVE" not in decision:
            doc.add_heading("1. Décision et synthèse", 1)
            fallback_synthese = (
                f"L'analyse de l'appel d'offres '{ao.titre}' de {ao.client} ne peut pas être "
                f"conclue de façon fiable en l'état — des éléments de configuration manquent "
                f"pour statuer avec certitude"
                + (f" — {', '.join(result.scoring_missing_display())}" if result.scoring_missing else "") + "."
                if decision == "INCOMPLET" else
                f"Après analyse, nous avons décidé de ne pas répondre à l'appel d'offres "
                f"'{ao.titre}' de {ao.client}. Les critères bloquants identifiés rendent "
                f"une réponse dans les conditions actuelles trop risquée."
            )
            doc.add_paragraph(ai.get("note_refus") or fallback_synthese)

            if result.criteres_bloquants:
                doc.add_heading("2. Critères bloquants", 1)
                doc.add_paragraph(ai.get("analyse_blocages") or "")
                for b in result.criteres_bloquants:
                    doc.add_paragraph(b, style="List Bullet")

            doc.add_heading("3. Risques si réponse quand même", 1)
            doc.add_paragraph(ai.get("risques_si_reponse") or
                "Répondre malgré les blocages exposerait notre structure à des risques opérationnels "
                "et réputationnels significatifs.")

            doc.add_heading("4. Conditions pour reconsidérer", 1)
            conditions = ai.get("conditions_reconsideration") or result.recommandations
            items = conditions if isinstance(conditions, list) else [conditions]
            for item in items:
                doc.add_paragraph(str(item), style="List Bullet")

            doc.add_heading("5. Alternatives et maintien de la relation", 1)
            doc.add_paragraph(ai.get("alternatives_proposees") or
                f"Nous recommandons de maintenir la relation avec {ao.client} et de "
                f"nous positionner sur les prochains marchés une fois les blocages levés.")

            doc.add_heading("6. Scoring détaillé (référence interne)", 1)
            for c in result.criteres:
                p = doc.add_paragraph()
                p.add_run(f"{c.nom} ({int(c.poids)}%) : ").bold = True
                p.add_run(_criterion_text(c))

        # ════════════════════════════════
        # CAS GO SOUS RÉSERVE
        # ════════════════════════════════
        elif "RESERVE" in decision:
            doc.add_paragraph(
                "⚠️  DOCUMENT CONDITIONNEL — Ce dossier est préparé sous réserve de validation "
                "des points listés en section 3. Ne pas soumettre avant levée des réserves.",
                style="Intense Quote" if "Intense Quote" in [s.name for s in doc.styles] else "Normal"
            )

            doc.add_heading("1. Résumé exécutif", 1)
            doc.add_paragraph(ai.get("resume_executif") or
                f"Cette candidature pour le projet '{ao.titre}' de {ao.client} est préparée "
                f"sous réserve de validation de points spécifiques détaillés ci-après.")

            doc.add_heading("2. Compréhension du besoin", 1)
            doc.add_paragraph(ai.get("comprehension_besoin") or
                f"Le projet de {ao.client} porte sur {', '.join(ao.technologies_demandees[:3]) or 'les technologies demandées'}.")
            if ao.contraintes:
                for c in ao.contraintes[:4]:
                    doc.add_paragraph(c, style="List Bullet")

            doc.add_heading("3. Réserves et conditions à lever", 1)
            doc.add_paragraph(ai.get("reserves_et_conditions") or
                "Les points suivants doivent être validés avant soumission définitive :")
            if result.criteres_bloquants:
                for b in result.criteres_bloquants:
                    doc.add_paragraph(f"🔴 {b}", style="List Bullet")
            # Lot 44: no hardcoded 55/70 band — the criteria the policy's own
            # display threshold flags as weak, plus every criterion that could
            # not be evaluated (a reserve to lift, not a score).
            faiblesses_notables = [c for c in result.criteres if _criterion_state(c) in ("evalue", "hypothese") and c.nom in (result.faiblesses or [])]
            for f in faiblesses_notables[:3]:
                doc.add_paragraph(f"🟡 {f.nom} : {_criterion_text(f)}", style="List Bullet")
            for f in [c for c in result.criteres if _criterion_state(c) in ("manquant", "non_applicable")][:5]:
                doc.add_paragraph(f"🟡 {f.nom} : {_criterion_text(f)}", style="List Bullet")

            doc.add_heading("4. Plan d'action pour GO définitif", 1)
            plan = ai.get("plan_action_go") or result.recommandations
            items = plan if isinstance(plan, list) else [plan]
            for item in items:
                doc.add_paragraph(str(item), style="List Bullet")

            doc.add_heading("5. Méthodologie proposée", 1)
            doc.add_paragraph(ai.get("methodologie") or
                "Approche par phases liées aux livrables attendus, avec points de validation client à chaque jalon.")

            doc.add_heading("6. Équipe projet", 1)
            doc.add_paragraph(ai.get("equipe_proposee") or
                f"Équipe constituée autour des compétences clés : {', '.join(ao.technologies_demandees[:3]) or 'demandées'}.")

            doc.add_heading("7. Références similaires", 1)
            if result.rag_synthesis:
                doc.add_paragraph(result.rag_synthesis)
            for ev in result.evidence_pack[:3]:
                p = doc.add_paragraph()
                p.add_run(f"{ev.source}  |  Pertinence : {min(int(ev.score*400),100)}%").bold = True
                doc.add_paragraph(ev.content[:400])

            valeur = ai.get("valeur_ajoutee")
            if valeur:
                doc.add_heading("8. Notre valeur ajoutée", 1)
                items = valeur if isinstance(valeur, list) else [valeur]
                for item in items:
                    doc.add_paragraph(str(item), style="List Bullet")

            doc.add_heading("9. Conclusion", 1)
            doc.add_paragraph(ai.get("conclusion") or
                f"Sous réserve de validation des points mentionnés, nous sommes prêts à nous engager "
                f"pleinement sur ce projet. Nous proposons un échange avec {ao.client} avant soumission.")

        # ════════════════════════════════
        # CAS GO
        # ════════════════════════════════
        else:
            doc.add_heading("1. Résumé exécutif", 1)
            doc.add_paragraph(ai.get("resume_executif") or
                f"Nous répondons à l'appel d'offres '{ao.titre}' émis par {ao.client} avec "
                f"un positionnement fort sur la stack demandée et des références sectorielles probantes.")

            doc.add_heading("2. Compréhension du besoin", 1)
            doc.add_paragraph(ai.get("comprehension_besoin") or
                f"Le projet de {ao.client} porte sur {', '.join(ao.technologies_demandees[:3]) or 'les technologies demandées'}.")
            if ao.contraintes:
                doc.add_paragraph("Contraintes identifiées :")
                for c in ao.contraintes[:5]:
                    doc.add_paragraph(c, style="List Bullet")

            doc.add_heading("3. Évaluation de l'adéquation", 1)
            # Lot 44: every evaluated criterion, without a hardcoded 78 cut-off
            # (a criterion not evaluated is listed as such, never scored).
            for c in result.criteres[:12]:
                if _criterion_state(c) == "non_applicable" and getattr(c, "motif", None) == "disabled_by_policy":
                    continue
                p = doc.add_paragraph()
                p.add_run(f"{'✓' if _criterion_state(c) in ('evalue', 'hypothese') else '•'} {c.nom} : ").bold = True
                p.add_run(_criterion_text(c))

            doc.add_heading("4. Méthodologie proposée", 1)
            doc.add_paragraph(ai.get("methodologie") or
                "Approche par phases liées aux livrables attendus, avec points de validation client à chaque jalon clé.")

            doc.add_heading("5. Équipe projet", 1)
            doc.add_paragraph(ai.get("equipe_proposee") or
                f"Équipe dédiée avec expertise sur {', '.join(ao.technologies_demandees[:3]) or 'les technologies demandées'}.")

            doc.add_heading("6. Références et expériences similaires", 1)
            if result.rag_synthesis:
                doc.add_paragraph(result.rag_synthesis)
            for ev in result.evidence_pack[:4]:
                p = doc.add_paragraph()
                p.add_run(f"{ev.source}  |  Pertinence : {min(int(ev.score*400),100)}%").bold = True
                doc.add_paragraph(ev.content[:500])

            valeur = ai.get("valeur_ajoutee")
            if valeur:
                doc.add_heading("7. Notre valeur ajoutée", 1)
                items = valeur if isinstance(valeur, list) else [valeur]
                for item in items:
                    doc.add_paragraph(str(item), style="List Bullet")

            doc.add_heading("8. Recommandations", 1)
            for r in result.recommandations:
                doc.add_paragraph(r, style="List Bullet")

            doc.add_heading("9. Conclusion", 1)
            doc.add_paragraph(ai.get("conclusion") or
                f"Nous sommes pleinement mobilisés pour répondre à ce projet de {ao.client}. "
                f"Nous sommes disponibles pour tout échange complémentaire.")

        doc.save(str(tmp_path))
        return _finalize(tmp_path, path)

    def generate_pdf(self, ao: AOContext, result: ScoringResult, llm=None, *, output_path: Optional[Path] = None) -> Path:
        if output_path is None:
            # Legacy flat naming — kept exactly as-is for callers that don't
            # scope their own path.
            safe = _safe_name(ao.titre)
            ts = datetime.now().strftime("%Y%m%d_%H%M")
            path = self.output_dir / f"rapport_decision_{safe}_{ts}.pdf"
        else:
            path = _prepare_target(Path(output_path))
        tmp_path = path.with_name(path.name + ".tmp")

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("CustomTitle", parent=styles["Title"], fontSize=18, spaceAfter=12)
        h2_style = ParagraphStyle("CustomH2", parent=styles["Heading2"], fontSize=13, spaceBefore=10)
        normal = styles["Normal"]

        # Decision color
        decision_color = colors.green if result.decision == "GO" else (
            colors.orange if "RESERVE" in result.decision else colors.red
        )
        decision_style = ParagraphStyle(
            "Decision", parent=styles["Heading1"],
            textColor=decision_color, fontSize=16
        )

        ai = result.ai_content or (self._generate_ai_content(ao, result, llm) if llm else {})

        story = [
            Paragraph("Rapport de décision — WinMarket AI", title_style),
            HRFlowable(width="100%", thickness=1, color=colors.grey),
            Spacer(1, 12),
            Paragraph(f"Appel d'offres : {_escape_for_pdf(ao.titre)}", h2_style),
            Paragraph(f"Client : {_escape_for_pdf(ao.client)}", normal),
            Paragraph(f"Secteur : {_escape_for_pdf(ao.secteur) or 'Non renseigné'}", normal),
            Spacer(1, 8),
            Paragraph(f"DÉCISION : {_escape_for_pdf(result.decision)}", decision_style),
            Paragraph(f"Score global : {_escape_for_pdf(_score_line(result))}", h2_style),
            Spacer(1, 12),
        ]

        # Lot 47 bis — pieces read, conflicts left unresolved, provenance of the facts
        d_pieces, d_conflicts, d_provenance, d_scope_note = _dossier_lines(ao)
        if d_pieces:
            story.append(Paragraph("Dossier d'appel d'offres analysé", h2_style))
            if d_scope_note:  # lot 50 §4/§5 — distinct from the decision text above, never merged with it
                story.append(Paragraph(f"<i>{_escape_for_pdf(d_scope_note)}</i>", normal))
            for line in d_pieces:
                story.append(Paragraph(f"• {_escape_for_pdf(line)}", normal))
            if d_conflicts:
                story.append(Paragraph("<b>Valeurs contradictoires entre les pièces (non tranchées)</b>", normal))
                for line in d_conflicts:
                    story.append(Paragraph(f"• {_escape_for_pdf(line)}", normal))
            if d_provenance:
                story.append(Paragraph("<b>Provenance des informations relevées</b>", normal))
                for line in d_provenance:
                    story.append(Paragraph(f"• {_escape_for_pdf(line)}", normal))
            story.append(Spacer(1, 8))

        # Lot 53 — this revision's own before/after (frozen at submission time, never recomputed here).
        change_lines = _completion_change_lines(result)
        if change_lines:
            story.append(Paragraph("Compléments et changements de cette révision", h2_style))
            for line in change_lines:
                story.append(Paragraph(f"• {_escape_for_pdf(line)}", normal))
            story.append(Spacer(1, 8))

        # Résumé exécutif IA
        resume = ai.get("resume_executif")
        if resume:
            story.append(Paragraph("Résumé exécutif", h2_style))
            story.append(Paragraph(_escape_for_pdf(resume), normal))
            story.append(Spacer(1, 8))

        # Synthèse RAG
        if result.rag_synthesis:
            story.append(Paragraph("Analyse des références internes", h2_style))
            story.append(Paragraph(_escape_for_pdf(result.rag_synthesis), normal))
            story.append(Spacer(1, 8))

        # Critères
        story.append(Paragraph("Évaluation détaillée par critère", h2_style))
        for c in result.criteres:
            state = _criterion_state(c)
            # The <b>…</b> here is markup this file deliberately emits; only
            # the interpolated criterion NAME is dynamic and escaped.
            if state in ("manquant", "non_applicable"):
                head = "NON ÉVALUÉ" if state == "manquant" else "NON APPLICABLE"
                story.append(Paragraph(f"<b>{_escape_for_pdf(c.nom)}</b> ({c.poids}%) : {head}", normal))
            else:
                bar = "█" * int(c.score / 10) + "░" * (10 - int(c.score / 10))
                mark = " (hypothèse de la politique)" if state == "hypothese" else ""
                story.append(Paragraph(
                    f"<b>{_escape_for_pdf(c.nom)}</b> ({c.poids}%) : {c.score:.0f}/100{mark}  {bar}",
                    normal
                ))
            story.append(Paragraph(f"  → {_escape_for_pdf(c.justification)}", normal))
            story.append(Spacer(1, 4))

        # Blockers
        if result.criteres_bloquants:
            story.append(Spacer(1, 8))
            story.append(Paragraph("Criteres bloquants", h2_style))
            for b in result.criteres_bloquants:
                story.append(Paragraph(
                    f"⛔ {_escape_for_pdf(b)}", ParagraphStyle("Blocker", parent=normal, textColor=colors.red)
                ))
            story.append(Spacer(1, 8))

        # Forces / faiblesses
        if result.forces:
            story.append(Paragraph("Points forts", h2_style))
            for f in result.forces:
                story.append(Paragraph(f"✓ {_escape_for_pdf(f)}", normal))
        if result.faiblesses:
            story.append(Paragraph("Points faibles", h2_style))
            for f in result.faiblesses:
                story.append(Paragraph(f"✗ {_escape_for_pdf(f)}", normal))

        # Recommandations
        story.append(Spacer(1, 8))
        story.append(Paragraph("Recommandations", h2_style))
        for r in result.recommandations:
            story.append(Paragraph(f"• {_escape_for_pdf(r)}", normal))

        # Valeur ajoutée IA
        valeur = ai.get("valeur_ajoutee")
        if valeur:
            story.append(Spacer(1, 8))
            story.append(Paragraph("Notre valeur ajoutée", h2_style))
            items = valeur if isinstance(valeur, list) else [valeur]
            for item in items:
                story.append(Paragraph(f"✓ {_escape_for_pdf(item)}", normal))

        # Conclusion IA
        conclusion = ai.get("conclusion")
        if conclusion:
            story.append(Spacer(1, 8))
            story.append(Paragraph("Conclusion", h2_style))
            story.append(Paragraph(_escape_for_pdf(conclusion), normal))

        doc = SimpleDocTemplate(str(tmp_path), pagesize=A4, rightMargin=2*cm, leftMargin=2*cm)
        doc.build(story)
        return _finalize(tmp_path, path)
